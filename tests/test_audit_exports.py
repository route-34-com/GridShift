import csv
import io

from openpyxl import load_workbook

from backend.services.exporter import safe_cell
from tests.conftest import PASSWORD, add_user, make_client, sign_in


def entries(client, **params):
    return client.get("/api/audit", params=params).json()


def test_audit_records_who_what_where(admin, settings):
    make_client(settings).post("/api/auth/login", json={"email": "admin@example.com", "password": "wrong-pass-1"}, headers={"user-agent": "Pytest Browser"})
    user_id = admin.post("/api/users/invite", json={"email": "pat@example.com", "role": "viewer", "name": "Pat"}).json()["id"]
    admin.put(f"/api/users/{user_id}", json={"role": "planner", "name": "Patricia"})
    found = entries(admin)["entries"]
    by_action = {e["action"]: e for e in found}
    failed = by_action["auth.login_failed"]
    assert failed["outcome"] == "failure" and failed["detail"]["reason"] == "wrong password"
    assert failed["user_agent"] == "Pytest Browser" and failed["ip"]
    change = by_action["user.update"]
    assert change["detail"]["before"] == {"name": "Pat", "role": "viewer"}
    assert change["detail"]["after"] == {"name": "Patricia", "role": "planner"}
    assert change["user_email"] == "admin@example.com" and change["user_name"] == "Ada Admin"
    assert by_action["user.invite"]["detail"]["email"] == "pat@example.com"


def test_audit_records_plan_runs_with_figures(admin):
    run = admin.post("/api/runs").json()
    entry = entries(admin, action="plan.run")["entries"][0]
    assert entry["entity_id"] == str(run["id"])
    assert entry["detail"]["trigger"] == "manual"
    assert entry["detail"]["savings_eur"] == run["summary"]["savings_eur"]
    assert "saved" in entry["summary"]


def test_audit_filters_and_paging(admin, settings):
    add_user(settings, "op@example.com", "viewer")
    sign_in(settings, "op@example.com")
    op_id = next(u["id"] for u in admin.get("/api/users").json() if u["email"] == "op@example.com")
    assert all(e["user_id"] == op_id for e in entries(admin, user_id=op_id)["entries"])
    assert all(e["action"].startswith("auth.") for e in entries(admin, category="auth")["entries"])
    assert all(e["outcome"] == "failure" for e in entries(admin, outcome="failure")["entries"])
    assert entries(admin, q="op@example.com")["total"] >= 1
    assert entries(admin, **{"from": "2999-01-01"})["total"] == 0
    page = entries(admin, limit=1, offset=0)
    assert len(page["entries"]) == 1 and page["total"] >= 2


def test_audit_action_catalog(admin):
    actions = {a["action"] for a in admin.get("/api/audit/actions").json()}
    assert {"auth.login", "user.invite", "plan.run", "export.download"} <= actions


def test_run_exports_csv_and_excel(admin):
    admin.post("/api/runs")
    csv_response = admin.get("/api/exports/run/schedule?format=csv")
    assert csv_response.status_code == 200
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert 'filename="gridshift-plan' in csv_response.headers["content-disposition"]
    text = csv_response.content.decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(text)))
    assert rows and {"machine", "start_local", "end_local", "reason"} <= rows[0].keys()
    report = admin.get("/api/exports/run/report?format=xlsx")
    book = load_workbook(io.BytesIO(report.content))
    assert book.sheetnames == ["Summary", "Daily", "Schedule", "Machines", "Hourly"]
    assert book["Hourly"].max_row == 169
    assert book["Summary"]["A1"].value == "Item"


def test_every_run_export_works(admin):
    admin.post("/api/runs")
    for name in ("summary", "daily", "schedule", "baseline-schedule", "hourly", "baseline-hourly", "machines"):
        for fmt in ("csv", "xlsx"):
            assert admin.get(f"/api/exports/run/{name}?format={fmt}").status_code == 200, (name, fmt)


def test_export_errors(admin):
    assert admin.get("/api/exports/run/report?format=xlsx").status_code == 404
    admin.post("/api/runs")
    assert admin.get("/api/exports/run/report?format=csv").status_code == 400
    assert admin.get("/api/exports/run/schedule?format=pdf").status_code == 400
    assert admin.get("/api/exports/run/nonsense?format=csv").status_code == 404
    assert admin.get("/api/exports/run/schedule?format=csv&run_id=999").status_code == 404


def test_exports_are_audited(admin):
    admin.post("/api/runs")
    admin.get("/api/exports/run/daily?format=csv")
    entry = entries(admin, action="export.download")["entries"][0]
    assert entry["detail"]["export"] == "daily" and entry["detail"]["format"] == "csv" and entry["detail"]["rows"] == 7


def test_audit_and_users_export(admin):
    admin.post("/api/users/invite", json={"email": "=cmd@example.com", "role": "viewer"})
    audit = admin.get("/api/exports/audit?format=csv&category=user")
    rows = list(csv.DictReader(io.StringIO(audit.content.decode("utf-8-sig"))))
    assert rows and all(r["action"].startswith("user.") for r in rows)
    users = load_workbook(io.BytesIO(admin.get("/api/exports/users?format=xlsx").content))["Users"]
    emails = [users.cell(row=r, column=2).value for r in range(2, users.max_row + 1)]
    assert "'=cmd@example.com" in emails


def test_formula_injection_is_neutralised():
    assert safe_cell("=HYPERLINK(1)") == "'=HYPERLINK(1)"
    assert safe_cell("@SUM(A1)") == "'@SUM(A1)"
    assert safe_cell(-12.5) == -12.5
    assert safe_cell("normal") == "normal"


def test_planner_can_export_but_not_audit(admin, settings):
    admin.post("/api/runs")
    add_user(settings, "planner@example.com", "planner")
    planner = sign_in(settings, "planner@example.com")
    assert planner.get("/api/exports/run/report").status_code == 200
    assert planner.get("/api/exports/users").status_code == 403
    assert make_client(settings).post("/api/auth/login", json={"email": "planner@example.com", "password": PASSWORD}).status_code == 200
